"""
Streaming readers for uploaded training files (Train Data Source).

Every supported format is read in chunks of rows, so a large file is never
held in memory in one piece:

- .csv              - ml_utils.iter_csv_chunks (text, types inferred downstream)
- .parquet          - pyarrow row-group batches; the row count is read from the
                      file's metadata first, so an over-limit file fails at once
- .xlsx / .xlsm     - openpyxl read-only mode, one sheet (sheetName or the first)
- .jsonl / .ndjson  - one JSON object per line; a first pass collects every
                      key, so rows with different keys line up in one table
- .json             - an array of objects (or {"<key>": [objects]}), loaded
                      whole and therefore capped at ML_EXTRACT_JSON_MAX_BYTES

Values are written to CSV as given - nothing is guessed or converted, so text
like "001" stays text. Each reader also records which kind of value each
column held; column_dtypes() turns that into the dtypes saved next to the
output CSV (see ml_utils.write_dtypes_sidecar), so Int64-with-missing-values,
booleans, datetimes and categories survive into the next node.

Each chunk is read in a worker thread, so a timeout cancels the extraction
between chunks instead of leaving a whole-file parse running.
"""

from __future__ import annotations

import asyncio
import base64
import csv
import datetime as dt
import decimal
import io
import json
import math
from pathlib import Path
from typing import Any, AsyncIterator, Dict, Iterator, List, Optional, Sequence, Set, Tuple

import numpy as np
import pandas as pd

from app.core.exceptions.error_messages import ErrorKey
from app.core.exceptions.exception_classes import AppException

RowChunk = Tuple[List[str], List[Tuple[Any, ...]]]

CSV_SUFFIXES = {".csv"}
PARQUET_SUFFIXES = {".parquet"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
JSON_LINES_SUFFIXES = {".jsonl", ".ndjson"}
JSON_SUFFIXES = {".json"}
SUPPORTED_TRAINING_FILE_SUFFIXES = (
    CSV_SUFFIXES | PARQUET_SUFFIXES | EXCEL_SUFFIXES | JSON_LINES_SUFFIXES | JSON_SUFFIXES
)
# Shown in errors and the upload dialog, in this order.
SUPPORTED_TRAINING_FILE_LABEL = ".csv, .parquet, .xlsx, .xlsm, .json, .jsonl, .ndjson"

# A .json array is parsed whole (no streaming JSON parser is installed), so it
# is capped well below the streamed formats; larger data belongs in .jsonl.
DEFAULT_JSON_MAX_BYTES = 512 * 1024 * 1024

# Lines scanned per worker-thread call during the JSON Lines key pass.
_JSONL_SCAN_LINES = 20_000


def training_file_suffix(path: Any) -> str:
    return Path(str(path)).suffix.lower()


def ensure_supported_training_file(path: Any) -> str:
    """Return the file's suffix, or raise a clear error for unsupported types."""
    suffix = training_file_suffix(path)
    if suffix in SUPPORTED_TRAINING_FILE_SUFFIXES:
        return suffix
    if suffix == ".xls":
        detail = (
            "Excel 97-2003 (.xls) files aren't supported. Open the file in Excel, "
            "save it as .xlsx, and upload it again."
        )
    else:
        detail = (
            f"'{suffix or 'no extension'}' files aren't supported as training data. "
            f"Upload one of: {SUPPORTED_TRAINING_FILE_LABEL}."
        )
    raise AppException(error_key=ErrorKey.ML_EXTRACT_FILE_UNSUPPORTED, error_detail=detail)


def _invalid(detail: str) -> AppException:
    return AppException(error_key=ErrorKey.ML_EXTRACT_FILE_INVALID, error_detail=detail)


def _too_many_rows(row_count: int, max_rows: int) -> AppException:
    return AppException(
        error_key=ErrorKey.ML_EXTRACT_LIMIT_EXCEEDED,
        error_detail=(
            f"The training file has {row_count:,} rows, which exceeds the limit of "
            f"{max_rows:,}. Use a smaller file, or ask an administrator to raise the row limit."
        ),
    )


def _not_utf8() -> AppException:
    return AppException(
        error_key=ErrorKey.ML_EXTRACT_FILE_ENCODING_INVALID,
        error_detail=(
            "The uploaded file could not be decoded as UTF-8. "
            "Save it with UTF-8 encoding and upload it again."
        ),
    )


def _unique_columns(columns: Sequence[Any], where: str) -> List[str]:
    names = ["" if c is None else str(c).strip() for c in columns]
    seen: Set[str] = set()
    for name in names:
        if name in seen:
            raise _invalid(
                f"The {where} has two columns named '{name}'. Give every column a unique name "
                "and upload the file again."
            )
        seen.add(name)
    return names


def csv_cell(value: Any) -> Any:
    """A value as written to the output CSV: missing -> None, nested -> JSON text."""
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, decimal.Decimal):
        return str(value)
    if isinstance(value, (bytes, bytearray)):
        return base64.b64encode(bytes(value)).decode("ascii")
    if isinstance(value, (list, tuple, dict, np.ndarray)):
        return json.dumps(value.tolist() if isinstance(value, np.ndarray) else value, default=str)
    return str(value)


class ColumnTypeTracker:
    """Records which kinds of value each column held, for the saved dtypes."""

    def __init__(self) -> None:
        self._kinds: Dict[str, Set[str]] = {}
        self._overrides: Dict[str, str] = {}

    def set_columns(self, columns: Sequence[str]) -> None:
        for column in columns:
            self._kinds.setdefault(column, set())

    def override(self, column: str, dtype: str) -> None:
        self._overrides[column] = dtype

    def observe(self, columns: Sequence[str], row: Sequence[Any]) -> None:
        for column, value in zip(columns, row):
            kind = self._kind(value)
            if kind:
                self._kinds[column].add(kind)

    @staticmethod
    def _kind(value: Any) -> Optional[str]:
        if value is None or value is pd.NA or value is pd.NaT:
            return None
        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, bool):
            return "bool"
        if isinstance(value, int):
            return "int"
        if isinstance(value, float):
            return None if math.isnan(value) else "float"
        if isinstance(value, (dt.datetime, dt.date)):
            return "datetime"
        return "other"

    def column_dtypes(self) -> Dict[str, str]:
        dtypes: Dict[str, str] = {}
        for column, kinds in self._kinds.items():
            if column in self._overrides:
                dtypes[column] = self._overrides[column]
            elif kinds == {"bool"}:
                dtypes[column] = "boolean"
            elif kinds == {"int"}:
                dtypes[column] = "Int64"
            elif kinds and kinds <= {"int", "float"}:
                dtypes[column] = "float64"
            elif kinds == {"datetime"}:
                dtypes[column] = "datetime64[ns]"
            else:
                dtypes[column] = "object"
        return dtypes


class TrainingFileReader:
    """Reads one uploaded training file as (columns, rows) chunks."""

    def __init__(
        self,
        path: Any,
        chunk_rows: int,
        *,
        max_rows: Optional[int] = None,
        sheet_name: Optional[str] = None,
        json_max_bytes: int = DEFAULT_JSON_MAX_BYTES,
        scan_rows: Optional[int] = None,
    ) -> None:
        if chunk_rows <= 0:
            raise ValueError("chunk_rows must be a positive integer")
        self.path = Path(str(path))
        self.suffix = ensure_supported_training_file(self.path)
        self.chunk_rows = chunk_rows
        self.max_rows = max_rows
        self.sheet_name = (sheet_name or "").strip() or None
        self.json_max_bytes = json_max_bytes
        # JSON Lines: stop the key pass after this many records (a preview
        # needs only the first rows, not the whole file's keys and count).
        self.scan_rows = scan_rows
        self.types = ColumnTypeTracker()
        # Known up front for Parquet (from metadata); None otherwise.
        self.total_rows: Optional[int] = None

    @property
    def is_csv(self) -> bool:
        return self.suffix in CSV_SUFFIXES

    def column_dtypes(self) -> Dict[str, str]:
        return self.types.column_dtypes()

    async def chunks(self) -> AsyncIterator[RowChunk]:
        """Yield (columns, []) once, then (columns, rows) chunks."""
        if self.is_csv:
            from app.modules.workflow.engine.nodes.ml import ml_utils

            async for chunk in ml_utils.iter_csv_chunks(str(self.path), self.chunk_rows):
                yield chunk
            return

        if self.suffix in PARQUET_SUFFIXES:
            source = await asyncio.to_thread(self._open_parquet)
        elif self.suffix in EXCEL_SUFFIXES:
            source = await asyncio.to_thread(self._open_excel)
        elif self.suffix in JSON_LINES_SUFFIXES:
            source = await self._open_json_lines()
        else:
            source = await asyncio.to_thread(self._open_json)

        columns, next_rows, close = source
        try:
            self.types.set_columns(columns)
            yield columns, []
            while True:
                rows = await asyncio.to_thread(next_rows)
                if rows is None:
                    break
                for row in rows:
                    self.types.observe(columns, row)
                yield columns, [tuple(csv_cell(v) for v in row) for row in rows]
        finally:
            close()

    # -- Parquet -----------------------------------------------------------

    def _open_parquet(self):
        import pyarrow as pa
        import pyarrow.parquet as pq

        try:
            parquet = pq.ParquetFile(str(self.path))
        except Exception as exc:
            raise _invalid(
                "The file isn't a readable Parquet file. Check that it was exported as "
                "Parquet and upload it again."
            ) from exc

        self.total_rows = parquet.metadata.num_rows
        if self.max_rows is not None and self.total_rows > self.max_rows:
            parquet.close()
            raise _too_many_rows(self.total_rows, self.max_rows)

        schema = parquet.schema_arrow
        columns = _unique_columns(schema.names, "Parquet file")
        for field, column in zip(schema, columns):
            if pa.types.is_dictionary(field.type):
                self.types.override(column, "category")
        batches = parquet.iter_batches(batch_size=self.chunk_rows)

        def next_rows():
            batch = next(batches, None)
            if batch is None:
                return None
            values = [batch.column(i).to_pylist() for i in range(batch.num_columns)]
            return [tuple(row) for row in zip(*values)] if values else []

        return columns, next_rows, parquet.close

    # -- Excel -------------------------------------------------------------

    def _open_excel(self):
        import zipfile

        from openpyxl import load_workbook

        try:
            workbook = load_workbook(str(self.path), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
            raise _invalid(
                "The file isn't a readable Excel workbook. Open it in Excel, save it as "
                ".xlsx, and upload it again."
            ) from exc
        except Exception as exc:  # openpyxl raises its own InvalidFileException
            raise _invalid(
                "The file isn't a readable Excel workbook. Open it in Excel, save it as "
                ".xlsx, and upload it again."
            ) from exc

        try:
            if self.sheet_name:
                if self.sheet_name not in workbook.sheetnames:
                    raise AppException(
                        error_key=ErrorKey.ML_EXTRACT_CONFIGURATION_INVALID,
                        error_detail=(
                            f"The workbook has no sheet named '{self.sheet_name}'. "
                            f"Sheets: {', '.join(workbook.sheetnames)}."
                        ),
                    )
                sheet = workbook[self.sheet_name]
            else:
                sheet = workbook.worksheets[0]
            rows_iter: Iterator[Tuple[Any, ...]] = sheet.iter_rows(values_only=True)

            header: Optional[Tuple[Any, ...]] = None
            header_line = 0
            for header_line, candidate in enumerate(rows_iter, start=1):
                if any(v is not None and str(v).strip() != "" for v in candidate):
                    header = candidate
                    break
            if header is None:
                raise _invalid(f"The sheet '{sheet.title}' is empty. Add a header row and data.")

            width = len(header)
            while width and (header[width - 1] is None or str(header[width - 1]).strip() == ""):
                width -= 1
            columns = [
                str(v).strip() if v is not None and str(v).strip() else f"column_{i + 1}"
                for i, v in enumerate(header[:width])
            ]
            columns = _unique_columns(columns, f"sheet '{sheet.title}'")
        except BaseException:
            workbook.close()
            raise

        line = header_line

        def next_rows():
            nonlocal line
            rows: List[Tuple[Any, ...]] = []
            for raw in rows_iter:
                line += 1
                if not any(v is not None and (not isinstance(v, str) or v.strip()) for v in raw):
                    continue  # blank row
                if any(v is not None and (not isinstance(v, str) or v.strip()) for v in raw[width:]):
                    raise _invalid(
                        f"Row {line} of sheet '{sheet.title}' has more values than the header has "
                        "columns. Add a header for every column."
                    )
                row = tuple(raw[:width]) + (None,) * max(0, width - len(raw))
                rows.append(row)
                if len(rows) >= self.chunk_rows:
                    return rows
            return rows or None

        return columns, next_rows, workbook.close

    # -- JSON Lines --------------------------------------------------------

    async def _open_json_lines(self):
        handle = open(self.path, "r", encoding="utf-8-sig")
        try:
            columns, row_count, complete = await self._scan_json_lines(handle)
            if self.max_rows is not None and row_count > self.max_rows:
                raise _too_many_rows(row_count, self.max_rows)
            self.total_rows = row_count if complete else None
            handle.seek(0)
        except BaseException:
            handle.close()
            raise

        line = 0

        def next_rows():
            nonlocal line
            rows: List[Tuple[Any, ...]] = []
            try:
                for text in handle:
                    line += 1
                    if not text.strip():
                        continue
                    record = json.loads(text)
                    rows.append(tuple(record.get(c) for c in columns))
                    if len(rows) >= self.chunk_rows:
                        return rows
            except UnicodeDecodeError:
                raise _not_utf8() from None
            return rows or None

        return columns, next_rows, handle.close

    async def _scan_json_lines(self, handle) -> Tuple[List[str], int, bool]:
        """First pass: every key (in first-seen order) and the row count.

        With scan_rows set, stops at the record after that many; the returned
        flag says whether the whole file was scanned (the count is exact).
        """
        keys: Dict[str, None] = {}
        state = {"line": 0, "rows": 0, "complete": True}

        def scan_some() -> bool:
            for _ in range(_JSONL_SCAN_LINES):
                try:
                    text = handle.readline()
                except UnicodeDecodeError:
                    raise _not_utf8() from None
                if not text:
                    return False
                state["line"] += 1
                if not text.strip():
                    continue
                if self.scan_rows is not None and state["rows"] >= self.scan_rows:
                    state["complete"] = False
                    return False
                try:
                    record = json.loads(text)
                except json.JSONDecodeError as exc:
                    raise _invalid(
                        f"Line {state['line']} isn't valid JSON ({exc.msg}). Each line of a JSON "
                        "Lines file must be one JSON object."
                    ) from None
                if not isinstance(record, dict):
                    raise _invalid(
                        f"Line {state['line']} is a JSON {type(record).__name__}, not an object. "
                        "Each line must be one object such as {\"price\": 10, \"size\": \"M\"}."
                    )
                for key in record:
                    keys.setdefault(str(key), None)
                state["rows"] += 1
                if self.max_rows is not None and state["rows"] > self.max_rows:
                    raise _too_many_rows(state["rows"], self.max_rows)
            return True

        while await asyncio.to_thread(scan_some):
            pass
        if not keys:
            raise _invalid("The JSON Lines file has no records.")
        return list(keys), state["rows"], state["complete"]

    # -- JSON array --------------------------------------------------------

    def _open_json(self):
        size = self.path.stat().st_size
        if size > self.json_max_bytes:
            raise AppException(
                error_key=ErrorKey.ML_EXTRACT_LIMIT_EXCEEDED,
                error_detail=(
                    f"The JSON file is {size:,} bytes; a .json array can be at most "
                    f"{self.json_max_bytes:,} bytes because it has to be read in one piece. "
                    "Save the data as JSON Lines (.jsonl, one object per line), which has no "
                    "such limit, and upload it again."
                ),
            )
        try:
            with open(self.path, "r", encoding="utf-8-sig") as handle:
                text = handle.read()
        except UnicodeDecodeError:
            raise _not_utf8() from None
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            if exc.msg == "Extra data":
                # One object per line saved as .json - read it as JSON Lines.
                return self._json_lines_from_text(text)
            raise _invalid(
                f"The file isn't valid JSON ({exc.msg} at line {exc.lineno}). "
                "Use an array of objects, e.g. [{\"price\": 10}, {\"price\": 12}]."
            ) from None
        del text

        if isinstance(data, dict):
            lists = [v for v in data.values() if isinstance(v, list)]
            if len(lists) == 1:
                data = lists[0]
        if not isinstance(data, list):
            raise _invalid(
                "The JSON file must hold an array of objects, e.g. "
                "[{\"price\": 10}, {\"price\": 12}], or one key holding such an array."
            )
        return self._records_source(data)

    def _json_lines_from_text(self, text: str):
        records = []
        for line, item in enumerate(text.splitlines(), start=1):
            if not item.strip():
                continue
            try:
                records.append(json.loads(item))
            except json.JSONDecodeError as exc:
                raise _invalid(f"Line {line} isn't valid JSON ({exc.msg}).") from None
        return self._records_source(records)

    def _records_source(self, records: List[Any]):
        keys: Dict[str, None] = {}
        for index, record in enumerate(records, start=1):
            if not isinstance(record, dict):
                raise _invalid(
                    f"Item {index} of the JSON array is a {type(record).__name__}, not an object. "
                    "Every item must be an object such as {\"price\": 10}."
                )
            for key in record:
                keys.setdefault(str(key), None)
        if not keys:
            raise _invalid("The JSON file has no records.")
        self.total_rows = len(records)
        if self.max_rows is not None and len(records) > self.max_rows:
            raise _too_many_rows(len(records), self.max_rows)
        columns = list(keys)
        position = 0

        def next_rows():
            nonlocal position
            if position >= len(records):
                return None
            batch = records[position:position + self.chunk_rows]
            position += len(batch)
            return [tuple(record.get(c) for c in columns) for record in batch]

        return columns, next_rows, lambda: None


def list_excel_sheets(path: Any) -> List[str]:
    """Sheet names of a workbook, in workbook order."""
    import zipfile

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(str(path), read_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise _invalid(
            "The file isn't a readable Excel workbook. Open it in Excel, save it as "
            ".xlsx, and upload it again."
        ) from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


async def read_training_file_preview(
    path: Any,
    max_rows: int,
    *,
    sheet_name: Optional[str] = None,
    chunk_rows: int = 5_000,
    json_max_bytes: int = DEFAULT_JSON_MAX_BYTES,
) -> Tuple[pd.DataFrame, Optional[int], bool]:
    """The first max_rows rows of a training file as a typed DataFrame.

    Returns (frame, total_rows, truncated). total_rows is exact when known
    (Parquet metadata, or the whole file fit in max_rows), otherwise None.
    """
    reader = TrainingFileReader(
        path,
        min(chunk_rows, max_rows),
        sheet_name=sheet_name,
        json_max_bytes=json_max_bytes,
        scan_rows=max_rows,
    )
    columns: List[str] = []
    rows: List[Tuple[Any, ...]] = []
    truncated = False
    chunks = reader.chunks()
    try:
        async for chunk_columns, chunk_rows_ in chunks:
            columns = chunk_columns
            room = max_rows - len(rows)
            if len(chunk_rows_) > room:
                rows.extend(chunk_rows_[:room])
                truncated = True
                break
            rows.extend(chunk_rows_)
            if len(rows) >= max_rows:
                # Exactly full: there may be more rows behind it.
                truncated = reader.total_rows is None or reader.total_rows > len(rows)
                break
    finally:
        await chunks.aclose()

    if reader.is_csv:
        frame = _csv_rows_frame(columns, rows)
    else:
        frame = apply_column_dtypes(pd.DataFrame(rows, columns=columns), reader.column_dtypes())
    total = reader.total_rows if reader.total_rows is not None else (None if truncated else len(rows))
    return frame, total, truncated


def _csv_rows_frame(columns: List[str], rows: List[Tuple[Any, ...]]) -> pd.DataFrame:
    """CSV rows (all text) typed the way pd.read_csv types the saved file.

    Re-parsing them gives the preview the same numeric/boolean columns the
    next node sees when it reads the extracted CSV.
    """
    if not columns:
        return pd.DataFrame(rows, columns=columns)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(columns)
    writer.writerows(rows)
    buffer.seek(0)
    return pd.read_csv(buffer)


def apply_column_dtypes(frame: pd.DataFrame, dtypes: Dict[str, str]) -> pd.DataFrame:
    """Best-effort cast of written-out values back to the recorded dtypes."""
    for column, dtype in dtypes.items():
        if column not in frame.columns or dtype == "object":
            continue
        try:
            if dtype.startswith("datetime64"):
                frame[column] = pd.to_datetime(frame[column], errors="coerce", format="mixed")
            else:
                frame[column] = frame[column].astype(dtype)
        except (TypeError, ValueError):
            pass
    return frame
